import json
import re
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from pathlib import Path

src = Path(r"C:\Users\fabio\Desktop\vinculacion denominadas.json")
dst = Path(r"C:\Users\fabio\Desktop\vinculacion_denominadas_agrupado.xlsx")

print("Leyendo JSON...")
with open(src, "r", encoding="utf-8") as f:
    data = json.load(f)

print(f"Total registros: {len(data)}")

# Agrupar por caja
groups = {}
for item in data:
    caja = item["caja"]
    groups.setdefault(caja, []).append(item)

print(f"Grupos (cajas distintas): {len(groups)}")

wb = Workbook()
# Eliminar hoja por defecto
wb.remove(wb.active)

fields = ["caja", "cuil_cuit", "denominacion", "id_Catastro_database",
          "id_cliente", "n_cuenta", "n_lote", "registro"]

def safe_sheet_name(name):
    name = name.strip()
    # Reemplazar caracteres inválidos para nombres de hoja en Excel
    name = re.sub(r'[\\/?*\[\]:]', '_', name)
    # Límite de 31 caracteres
    if len(name) > 31:
        name = name[:31]
    if not name:
        name = "Sheet"
    return name

# Ordenar grupos alfabéticamente
sorted_cajas = sorted(groups.keys(), key=lambda x: (not x.startswith("CAJA"), x))

for caja in sorted_cajas:
    items = groups[caja]
    sheet_name = safe_sheet_name(caja)
    # Manejar nombres duplicados
    if sheet_name in wb.sheetnames:
        i = 2
        while f"{sheet_name[:27]}_{i}" in wb.sheetnames:
            i += 1
        sheet_name = f"{sheet_name[:27]}_{i}"

    ws = wb.create_sheet(title=sheet_name)

    # Encabezados
    ws.append(fields)

    # Datos
    for item in items:
        ws.append([item.get(f, "") for f in fields])

    # Autoajustar ancho de columnas
    for col_idx, field in enumerate(fields, 1):
        max_len = len(str(field))
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx,
                                 min_row=2, max_row=min(len(items) + 1, 100)):
            for cell in row:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 50)

    print(f"  {sheet_name}: {len(items)} registros")

wb.save(dst)
print(f"\nArchivo guardado: {dst}")
print("Listo!")